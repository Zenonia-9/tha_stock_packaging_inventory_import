import base64
import io
import math
import uuid
import zipfile

import openpyxl
import xlsxwriter

from odoo import _, api, fields, models
from odoo.exceptions import UserError
from odoo.tools.float_utils import float_round

from .common import (
    MAX_FILE_BYTES, MAX_ROWS, check_quant, conversion, identity, internal,
    require_internal, require_manager,
)


class InventoryCountTemplate(models.Model):
    _name = 'tha.inventory.count.template'
    _description = 'Packaging inventory count template snapshot'
    _order = 'id desc'

    name = fields.Char(required=True, readonly=True)
    token = fields.Char(required=True, readonly=True, copy=False, index=True)
    company_id = fields.Many2one('res.company', required=True, readonly=True)
    snapshot = fields.Json(required=True, readonly=True)
    file_data = fields.Binary(attachment=True, readonly=True)
    filename = fields.Char(readonly=True)
    _token_unique = models.Constraint('UNIQUE(token)', 'A count template token must be unique.')

    @api.model_create_multi
    def create(self, vals_list):
        require_internal(self.env)
        return super().create(vals_list)

    def write(self, vals):
        require_internal(self.env)
        return super().write(vals)

    def unlink(self):
        require_internal(self.env)
        return super().unlink()

    @api.model
    def _export_quants(self, quants):
        require_manager(self.env)
        quants = quants.exists()
        if not quants or len(quants) > MAX_ROWS:
            raise UserError(_('Select between 1 and %s existing inventory lines.', MAX_ROWS))
        quants.check_access('read')
        rows = {}
        max_slots = 0
        conversion_errors = {}
        for quant in quants.sorted(lambda q: (q.location_id.display_name, q.product_id.display_name, q.id)):
            check_quant(quant)
            product = quant.product_id
            unit = product.uom_id
            try:
                conversion(unit, unit, product)
            except UserError as error:
                conversion_errors.setdefault((product.id, unit.id), str(error))
            slots = []
            for packaging in product.uom_ids.filtered('active') - unit:
                try:
                    factor = conversion(packaging, unit, product)
                except UserError as error:
                    conversion_errors.setdefault((product.id, packaging.id), str(error))
                    continue
                slots.append({
                    'uom_id': packaging.id, 'name': packaging.name, 'factor': factor,
                    'label': '%s\n(contains %s %s)' % (packaging.name, format(factor, '.12g'), unit.name),
                })
            slots.sort(key=lambda slot: (slot['factor'], slot['name'], slot['uom_id']))
            max_slots = max(max_slots, len(slots))
            rows[str(quant.id)] = {
                'identity': identity(quant),
                'labels': [product.default_code or '', product.display_name,
                           quant.location_id.complete_name, quant.lot_id.name or '',
                           quant.package_id.name or '', quant.owner_id.display_name or '', unit.name],
                'slots': slots,
                'inventory_factor': unit.factor,
                'rounding': unit.rounding,
            }
        if conversion_errors:
            raise UserError('\n\n'.join(conversion_errors.values()))
        snapshot = {'version': 1, 'max_slots': max_slots, 'rows': rows}
        template = internal(self).create({
            'name': _('Packaging count sheet — %s', fields.Date.today()),
            'token': uuid.uuid4().hex,
            'company_id': self.env.company.id,
            'snapshot': snapshot,
            'filename': 'Packaging Counts.xlsx',
        })
        template.write({'file_data': base64.b64encode(template._workbook())})
        return template

    def _headers(self, max_slots=None):
        self.ensure_one()
        if max_slots is None:
            max_slots = self.snapshot['max_slots']
        headers = ['Internal Reference', 'Product', 'Location', 'Inventory Unit', 'Loose Count']
        for index in range(max_slots):
            headers.extend(['Packaging %s' % (index + 1), 'Count %s' % (index + 1)])
        headers.extend(['Lot / Serial', 'Stock Package', 'Owner'])
        headers.extend(['Stock Line ID', 'Product ID', 'Location ID', 'Lot ID',
                        'Stock Package ID', 'Owner ID', 'Inventory Unit ID'])
        headers.extend('Packaging %s ID' % (index + 1) for index in range(max_slots))
        return headers

    def _row_values(self, row, max_slots=None):
        if max_slots is None:
            max_slots = self.snapshot['max_slots']
        labels = row['labels']
        values = labels[:3] + [labels[6], None]
        for index in range(max_slots):
            slot = row['slots'][index] if index < len(row['slots']) else None
            values.extend([slot['label'] if slot else '', None])
        values.extend(labels[3:6])
        ident = row['identity']
        values.extend(ident[key] for key in ('quant_id', 'product_id', 'location_id', 'lot_id',
                                            'package_id', 'owner_id', 'inventory_uom_id'))
        values.extend(row['slots'][index]['uom_id'] if index < len(row['slots']) else 0
                      for index in range(max_slots))
        return values

    def _workbook(self):
        self.ensure_one()
        # Json fields copy their value on access; read the large snapshot once.
        snapshot = self.snapshot
        max_slots = snapshot['max_slots']
        output = io.BytesIO()
        book = xlsxwriter.Workbook(output, {'in_memory': True, 'strings_to_formulas': False,
                                           'strings_to_urls': False})
        sheet = book.add_worksheet('Counts')
        header = book.add_format({'bold': True, 'bg_color': '#E7E6E6', 'text_wrap': True})
        label = book.add_format({'text_wrap': True, 'valign': 'vcenter', 'locked': True})
        count = book.add_format({'locked': False, 'bg_color': '#FFF2CC', 'num_format': '0.########'})
        unused = book.add_format({'bg_color': '#E7E6E6', 'locked': True})
        headers = self._headers(max_slots)
        visible_columns = 8 + max_slots * 2
        sheet.write_row(0, 0, headers, header)
        sheet.set_row(0, 30)
        sheet.set_column(0, 0, 18)
        sheet.set_column(1, 1, 45)
        sheet.set_column(2, 2, 30)
        sheet.set_column(3, 3, 18)
        sheet.set_column(4, 4, 14)
        for slot in range(max_slots):
            sheet.set_column(5 + slot * 2, 5 + slot * 2, 25)
            sheet.set_column(6 + slot * 2, 6 + slot * 2, 14)
        detail_start = 5 + max_slots * 2
        sheet.set_column(detail_start, detail_start + 2, 20)
        sheet.set_column(visible_columns, len(headers) - 1, 12, None, {'hidden': True})
        for row_index, row in enumerate(snapshot['rows'].values(), start=1):
            sheet.set_row(row_index, 34)
            sheet.write_row(row_index, 0, self._row_values(row, max_slots), label)
            for column in [4] + [6 + index * 2 for index in range(max_slots)]:
                is_unused = column != 4 and (column - 6) // 2 >= len(row['slots'])
                sheet.write_blank(row_index, column, None, unused if is_unused else count)
                if not is_unused:
                    sheet.data_validation(row_index, column, row_index, column, {
                        'validate': 'decimal', 'criteria': '>=', 'value': 0,
                        'input_title': 'Physical count', 'input_message': 'Blank = not counted; 0 = counted as zero.',
                        'error_title': 'Invalid count', 'error_message': 'Enter a nonnegative number.',
                    })
        sheet.freeze_panes(1, 2)
        sheet.autofilter(0, 0, len(snapshot['rows']), visible_columns - 1)
        sheet.protect('', {'autofilter': True, 'format_columns': True, 'format_rows': True,
                           'select_locked_cells': True, 'select_unlocked_cells': True})
        guide = book.add_worksheet('Instructions')
        guide.set_column(0, 0, 110)
        for index, text in enumerate([
            'Enter counts only in the yellow cells. Packaging names and sizes are already filled in.',
            'Use one row per existing inventory line. Do not copy rows or change product/location/unit labels.',
            'Leave every count blank to skip a line. On a counted line, blank packaging counts mean zero.',
            'Enter 0 in Loose Count to record a zero physical count. Do not enter quantities in gray cells.',
            'Upload this file using Upload Packaging Counts in Physical Inventory. Review before confirming.',
            'Confirmation updates Counted Quantity only. Use the normal Apply action to adjust stock.',
            'Download a fresh sheet after products or packaging conversions change.',
        ]):
            guide.write_string(index, 0, text)
        metadata = book.add_worksheet('_InventoryMetadata')
        metadata.write_row(0, 0, ['Version', 1])
        metadata.write_row(1, 0, ['Template Token', self.token])
        metadata.very_hidden()
        metadata.protect()
        book.close()
        return output.getvalue()

    def download_action(self):
        self.ensure_one()
        self.check_access('read')
        return {'type': 'ir.actions.act_url', 'target': 'download',
                'url': '/web/content/tha.inventory.count.template/%s/file_data/Packaging%%20Counts.xlsx?download=true' % self.id}

    def _check_live(self, row, counts):
        quant = self.env['stock.quant'].browse(row['identity']['quant_id'])
        check_quant(quant, row['identity'])
        unit = quant.product_id.uom_id
        if not unit.active or not math.isclose(unit.factor, row['inventory_factor'], rel_tol=1e-12) or unit.rounding != row['rounding']:
            raise UserError(_('The inventory unit or its precision changed. Download a fresh template.'))
        allowed = quant.product_id.uom_ids
        converted = []
        for entry in counts:
            packaging = self.env['uom.uom'].browse(entry['uom_id']).exists()
            if not packaging or (packaging != unit and packaging not in allowed):
                raise UserError(_('A counted packaging was removed from this product. Download a fresh template.'))
            factor = conversion(packaging, unit, quant.product_id)
            if not math.isclose(factor, entry['factor'], rel_tol=1e-12, abs_tol=1e-12):
                raise UserError(_('A counted packaging conversion changed. Download a fresh template.'))
            converted.append(packaging._compute_quantity(entry['quantity'], unit, round=False))
        try:
            total = math.fsum(converted)
        except OverflowError as error:
            raise UserError(_('The converted count is too large.')) from error
        if not math.isfinite(total):
            raise UserError(_('The converted count is too large.'))
        total = float_round(total, precision_rounding=unit.rounding, rounding_method='HALF-UP')
        if quant.product_id.tracking != 'none' and not quant.lot_id:
            raise UserError(_('A counted tracked product requires an existing lot or serial number on the inventory line.'))
        if quant.product_id.tracking == 'serial' and total not in (0.0, 1.0):
            raise UserError(_('A serial-number inventory line can only be counted as zero or one inventory unit.'))
        return quant, total

    @api.model
    def _read_upload(self, data):
        require_manager(self.env)
        if not data or len(data) > MAX_FILE_BYTES:
            raise UserError(_('Upload an .xlsx counting file no larger than 10 MB.'))
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 1000 or sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                    raise UserError(_('The workbook is too large to process.'))
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        except (zipfile.BadZipFile, OSError, ValueError, KeyError) as error:
            raise UserError(_('Upload the original .xlsx template downloaded from Physical Inventory.')) from error
        try:
            if '_InventoryMetadata' not in book or 'Counts' not in book:
                raise UserError(_('The counting template metadata or Counts sheet is missing.'))
            meta = book['_InventoryMetadata']
            if meta['B1'].value != 1 or not isinstance(meta['B2'].value, str):
                raise UserError(_('This counting template version is not supported.'))
            template = self.search([('token', '=', meta['B2'].value), ('company_id', '=', self.env.company.id)], limit=1)
            if not template:
                raise UserError(_('This template does not belong to the current company or database. Download a fresh sheet.'))
            template.check_access('read')
            snapshot = template.snapshot
            max_slots = snapshot['max_slots']
            expected_headers = template._headers(max_slots)
            sheet = book['Counts']
            if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > len(expected_headers):
                raise UserError(_('The workbook has too many rows or unexpected columns.'))
            headers = next(sheet.iter_rows(min_row=1, max_row=1, max_col=len(expected_headers), values_only=True))
            if list(headers) != expected_headers:
                raise UserError(_('The counting sheet headers were changed. Download a fresh template.'))
            lines, seen = [], set()
            hidden_start = 8 + max_slots * 2
            count_columns = {4} | {6 + index * 2 for index in range(max_slots)}
            for excel_row, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers)), start=2):
                if excel_row > MAX_ROWS + 1:
                    raise UserError(_('The workbook has too many rows.'))
                values = [cell.value for cell in cells]
                if all(value is None for value in values):
                    continue
                quant_id = values[hidden_start]
                if isinstance(quant_id, bool) or not isinstance(quant_id, (int, float)) or not math.isfinite(quant_id) or quant_id != int(quant_id):
                    raise UserError(_('Row %s: the stock-line ID is missing or invalid.', excel_row))
                quant_id = int(quant_id)
                if quant_id in seen:
                    raise UserError(_('Row %s: this stock-line ID appears more than once.', excel_row))
                seen.add(quant_id)
                row = snapshot['rows'].get(str(quant_id))
                if not row:
                    raise UserError(_('Row %s: this inventory line was not in the exported template.', excel_row))
                expected = template._row_values(row, max_slots)
                if any((value if value is not None else '') != (expected[index] if expected[index] is not None else '')
                       for index, value in enumerate(values) if index not in count_columns):
                    raise UserError(_('Row %s: product, location, packaging labels or hidden identities were changed.', excel_row))
                counts = []
                for column in sorted(count_columns):
                    value = values[column]
                    if value is None or value == '':
                        continue
                    if column == 4:
                        entry = {'uom_id': row['identity']['inventory_uom_id'], 'factor': 1.0, 'label': row['labels'][6]}
                    else:
                        index = (column - 6) // 2
                        if index >= len(row['slots']):
                            raise UserError(_('Row %s: a count was entered in an unused packaging slot.', excel_row))
                        entry = dict(row['slots'][index])
                    if cells[column].data_type == 'f' or isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
                        raise UserError(_('Row %s: enter a nonnegative finite number, not text or a formula.', excel_row))
                    entry['quantity'] = float(value)
                    counts.append(entry)
                if not counts:
                    continue
                try:
                    quant, total = template._check_live(row, counts)
                except UserError as error:
                    raise UserError(_('Row %(row)s: %(error)s', row=excel_row, error=str(error))) from error
                lines.append({
                    'stock_line_id': quant.id, 'quant_id': quant.id, 'source_row': excel_row,
                    'product_name': row['labels'][1], 'location_name': row['labels'][2],
                    'lot_name': row['labels'][3], 'package_name': row['labels'][4], 'owner_name': row['labels'][5],
                    'unit_name': row['labels'][6], 'counts': counts,
                    'old_count': quant.inventory_quantity, 'old_count_set': quant.inventory_quantity_set,
                    'on_hand': quant.quantity, 'new_count': total,
                    'count_details': '\n'.join('%s × %s = %s %s' % (
                        format(entry['quantity'], '.12g'), entry['label'].replace('\n', ' '),
                        format(entry['quantity'] * entry['factor'], '.12g'), row['labels'][6]) for entry in counts),
                })
            if not lines:
                raise UserError(_('No counts were entered. Enter at least one number, including zero.'))
            return template, lines
        finally:
            book.close()
