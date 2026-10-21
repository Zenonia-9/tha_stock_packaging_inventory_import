import base64
import io

import openpyxl

from odoo import Command
from odoo.exceptions import AccessError, UserError
from odoo.service.model import call_kw
from odoo.tests import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestPackagingInventory(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.unit = cls.env['uom.uom'].create({'name': 'Count ml', 'rounding': 0.01})
        cls.bottle = cls.env['uom.uom'].create({'name': 'Bottle', 'relative_uom_id': cls.unit.id, 'relative_factor': 250})
        cls.box = cls.env['uom.uom'].create({'name': 'Box', 'relative_uom_id': cls.bottle.id, 'relative_factor': 5})
        cls.sprite_box = cls.env['uom.uom'].create({'name': 'Box', 'relative_uom_id': cls.unit.id, 'relative_factor': 1000})
        cls.location = cls.env['stock.location'].create({
            'name': 'Packaging Count Test', 'usage': 'internal', 'company_id': cls.env.company.id,
        })
        cls.coke = cls.env['product.product'].create({
            'name': 'Count Coke', 'default_code': 'COUNT-COKE', 'is_storable': True,
            'uom_id': cls.unit.id, 'uom_ids': [Command.set([cls.box.id, cls.bottle.id])],
        })
        cls.sprite = cls.env['product.product'].create({
            'name': 'Count Sprite', 'default_code': 'COUNT-SPRITE', 'is_storable': True,
            'uom_id': cls.unit.id, 'uom_ids': [Command.set([cls.sprite_box.id, cls.bottle.id])],
        })
        cls.quants = cls.env['stock.quant'].create([
            {'product_id': cls.coke.id, 'location_id': cls.location.id, 'quantity': 5},
            {'product_id': cls.sprite.id, 'location_id': cls.location.id, 'quantity': 7},
        ])
        cls.templates = cls.env['tha.inventory.count.template']
        cls.batches = cls.env['tha.inventory.count.import']

    def _template(self, quants=None):
        return self.templates._export_quants(self.quants if quants is None else quants)

    def _edit_file(self, template, edits=None, mutate=None):
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(template.file_data)))
        sheet = book['Counts']
        hidden = 8 + template.snapshot['max_slots'] * 2
        for row in range(2, sheet.max_row + 1):
            quant_id = sheet.cell(row, hidden + 1).value
            for column, value in (edits or {}).get(quant_id, {}).items():
                sheet.cell(row, column, value)
        if mutate:
            mutate(book)
        output = io.BytesIO()
        book.save(output)
        book.close()
        return output.getvalue()

    def _preview(self, template, edits=None, mutate=None):
        return self.batches._preview_upload(self._edit_file(template, edits, mutate), 'Counts.xlsx')

    def test_mixed_counts_update_existing_without_moves(self):
        template = self._template()
        before_quants = self.env['stock.quant'].search_count([])
        before_moves = self.env['stock.move'].search_count([])
        before_valuation = self.env['account.move'].search_count([]) if 'account.move' in self.env else None
        quantities = self.quants.mapped('quantity')
        batch = self._preview(template, {
            self.quants[0].id: {5: 100, 7: 3, 9: 2},
            self.quants[1].id: {5: 0, 7: 1, 9: 2},
        })
        self.assertEqual(batch.line_ids.mapped('new_count'), [3350, 2250])
        self.assertEqual(self.quants.mapped('inventory_quantity'), [0, 0])
        batch.action_confirm()
        self.assertEqual(self.quants.mapped('inventory_quantity'), [3350, 2250])
        self.assertEqual(self.quants.mapped('quantity'), quantities)
        self.assertTrue(all(self.quants.mapped('inventory_quantity_set')))
        self.assertEqual(self.env['stock.quant'].search_count([]), before_quants)
        self.assertEqual(self.env['stock.move'].search_count([]), before_moves)
        if before_valuation is not None:
            self.assertEqual(self.env['account.move'].search_count([]), before_valuation)
        self.assertEqual(batch.state, 'done')
        self.assertTrue(batch.file_data)
        self.assertEqual(batch.imported_by, self.env.user)
        with self.assertRaises(UserError):
            batch.action_confirm()
        second = self._preview(template, {self.quants[0].id: {5: 100, 7: 3, 9: 2}})
        second.action_confirm()
        self.assertEqual(self.quants[0].inventory_quantity, 3350)

    def test_zero_blank_and_omitted_rows(self):
        self.quants.write({'inventory_quantity': 15, 'inventory_quantity_set': True})
        template = self._template()
        batch = self._preview(template, {self.quants[0].id: {5: 0}})
        self.assertEqual(len(batch.line_ids), 1)
        batch.action_confirm()
        self.assertEqual(self.quants.mapped('inventory_quantity'), [0, 15])
        with self.assertRaisesRegex(UserError, 'No counts'):
            self._preview(template)

    def test_invalid_quantities_and_unused_slots(self):
        plain = self.env['product.product'].create({'name': 'Plain count', 'is_storable': True, 'uom_id': self.unit.id})
        quant = self.env['stock.quant'].create({'product_id': plain.id, 'location_id': self.location.id})
        template = self._template(self.quants | quant)
        for value in (-1, 'two', True, '=1+1'):
            with self.subTest(value=value), self.assertRaises(UserError):
                self._preview(template, {self.quants[0].id: {5: value}})
        with self.assertRaisesRegex(UserError, 'unused packaging'):
            self._preview(template, {quant.id: {7: 0}})
        with self.assertRaisesRegex(UserError, 'too large'):
            self._preview(template, {self.quants[0].id: {7: 1e308, 9: 1e308}})

    def test_modified_identities_headers_and_duplicate_rows(self):
        template = self._template()
        for column in (1, 2, 3, 4, 6, 8, 10, 11, 12, 13, 14, 15, 19, 20):
            with self.subTest(column=column), self.assertRaises(UserError):
                self._preview(template, {self.quants[0].id: {5: 1}},
                              lambda book, col=column: book['Counts'].cell(2, col, 'Changed'))
        def duplicate(book):
            sheet = book['Counts']
            sheet.append([cell.value for cell in sheet[2]])
        with self.assertRaisesRegex(UserError, 'more than once'):
            self._preview(template, {self.quants[0].id: {5: 1}}, duplicate)
        with self.assertRaisesRegex(UserError, 'headers'):
            self._preview(template, {self.quants[0].id: {5: 1}}, lambda book: book['Counts'].cell(1, 1, 'Changed'))
        with self.assertRaises(UserError):
            self.batches._preview_upload(b'not an xlsx', 'Counts.xlsx')

    def test_packaging_conversion_changes_at_upload_and_confirmation(self):
        template = self._template()
        batch = self._preview(template, {self.quants[0].id: {9: 2}})
        # Changing an ancestor changes the box conversion as well.
        self.bottle.relative_factor = 200
        with self.assertRaisesRegex(UserError, 'conversion changed'):
            self._preview(template, {self.quants[0].id: {9: 2}})
        with self.assertRaisesRegex(UserError, 'conversion changed'):
            batch.action_confirm()
        self.assertEqual(self.quants[0].inventory_quantity, 0)

    def test_export_lists_all_products_with_invalid_packaging(self):
        unrelated = self.env['uom.uom'].create({'name': 'Unrelated packaging'})
        (self.coke | self.sprite).write({'uom_ids': [Command.link(unrelated.id)]})
        extra_quant = self.env['stock.quant'].create({
            'product_id': self.coke.id,
            'location_id': self.location.copy({'name': 'Other invalid packaging location'}).id,
        })
        before_templates = self.templates.search_count([])
        with self.assertRaises(UserError) as caught:
            self._template(self.quants | extra_quant)
        message = str(caught.exception)
        self.assertIn('same reference unit', message)
        for product in self.coke | self.sprite:
            self.assertEqual(message.count('Product: %s' % product.display_name), 1)
            self.assertIn(product.default_code, message)
        self.assertIn('Packaging: %s' % unrelated.display_name, message)
        self.assertIn('Inventory unit: %s' % self.unit.display_name, message)
        self.assertEqual(self.templates.search_count([]), before_templates)

    def test_mismatched_packaging_error_identifies_product_on_upload_and_confirmation(self):
        template = self._template()
        edits = {self.quants[0].id: {9: 2}}
        batch = self._preview(template, edits)
        self.box.relative_uom_id = self.env['uom.uom'].create({'name': 'Other reference unit'})
        for action in (lambda: self._preview(template, edits), batch.action_confirm):
            with self.assertRaises(UserError) as caught:
                action()
            message = str(caught.exception)
            self.assertIn('same reference unit', message)
            self.assertIn('Product: %s' % self.coke.display_name, message)
            self.assertIn('Packaging: %s' % self.box.display_name, message)
            self.assertIn('Inventory unit: %s' % self.unit.display_name, message)
        self.assertEqual(self.quants[0].inventory_quantity, 0)
        self.assertEqual(batch.state, 'preview')

    def test_packaging_removed_and_new_choices_refresh(self):
        template = self._template()
        self.coke.uom_ids = self.bottle
        with self.assertRaisesRegex(UserError, 'removed'):
            self._preview(template, {self.quants[0].id: {9: 1}})
        batch = self._preview(template, {self.quants[0].id: {5: 10}})
        batch.action_confirm()  # An unused removed slot does not prevent a loose count.
        fresh = self._template()
        self.assertEqual(len(fresh.snapshot['rows'][str(self.quants[0].id)]['slots']), 1)

    def test_changed_counts_after_preview_and_atomic_validation(self):
        template = self._template()
        batch = self._preview(template, {quant.id: {5: 10} for quant in self.quants})
        self.quants[1].inventory_quantity = 20
        with self.assertRaisesRegex(UserError, 'after the preview'):
            batch.action_confirm()
        self.assertEqual(self.quants[0].inventory_quantity, 0)
        self.assertEqual(batch.state, 'preview')

    def test_missing_quant_changed_dimension_and_inventory_precision(self):
        template = self._template()
        self.quants[0].unlink()
        with self.assertRaisesRegex(UserError, 'no longer exists'):
            self._preview(template, {self.quants[0].id: {5: 1}})
        template = self._template(self.quants[1])
        self.unit.rounding = 0.1
        with self.assertRaisesRegex(UserError, 'precision changed'):
            self._preview(template, {self.quants[1].id: {5: 1}})

    def test_round_after_sum_and_label_rename(self):
        template = self._template()
        self.bottle.name = 'Renamed Bottle'
        batch = self._preview(template, {self.quants[0].id: {5: 0.004, 7: 0.000016}})
        self.assertEqual(batch.line_ids.new_count, 0.01)
        batch.action_confirm()

    def test_lots_packages_owners_locations_and_serials(self):
        self.coke.tracking = 'lot'
        lot1, lot2 = self.env['stock.lot'].create([
            {'name': 'Count Lot A', 'product_id': self.coke.id, 'company_id': self.env.company.id},
            {'name': 'Count Lot B', 'product_id': self.coke.id, 'company_id': self.env.company.id},
        ])
        package = self.env['stock.package'].create({'name': 'Count Stock Package'})
        owner = self.env['res.partner'].create({'name': 'Count Owner'})
        other_location = self.location.copy({'name': 'Other Count Location'})
        quants = self.env['stock.quant'].create([
            {'product_id': self.coke.id, 'location_id': self.location.id, 'lot_id': lot1.id,
             'package_id': package.id, 'owner_id': owner.id, 'quantity': 3},
            {'product_id': self.coke.id, 'location_id': other_location.id, 'lot_id': lot2.id, 'quantity': 4},
        ])
        batch = self._preview(self._template(quants), {q.id: {5: 2} for q in quants})
        batch.action_confirm()
        self.assertEqual(quants.mapped('inventory_quantity'), [2, 2])
        self.assertEqual(quants.mapped('quantity'), [3, 4])
        self.assertEqual(quants[0].package_id, package)
        self.assertEqual(quants[0].owner_id, owner)
        with self.assertRaisesRegex(UserError, 'requires an existing lot'):
            self._preview(self._template(self.quants[0]), {self.quants[0].id: {5: 1}})
        self.sprite.tracking = 'serial'
        serial = self.env['stock.lot'].create({'name': 'Count Serial', 'product_id': self.sprite.id, 'company_id': self.env.company.id})
        serial_quant = self.env['stock.quant'].create({'product_id': self.sprite.id, 'location_id': self.location.id, 'lot_id': serial.id})
        serial_template = self._template(serial_quant)
        with self.assertRaisesRegex(UserError, 'zero or one'):
            self._preview(serial_template, {serial_quant.id: {7: 1}})
        self._preview(serial_template, {serial_quant.id: {5: 1}}).action_confirm()

    def test_workbook_protection_and_formula_free_layout(self):
        template = self._template()
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(template.file_data)))
        sheet = book['Counts']
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet.max_column, 21)
        self.assertTrue(sheet.protection.sheet)
        self.assertTrue(sheet.protection.formatRows)
        self.assertTrue(sheet.protection.formatColumns)
        self.assertFalse(sheet['E2'].protection.locked)
        self.assertTrue(sheet['B2'].protection.locked)
        self.assertEqual([sheet.cell(1, column).value for column in range(1, 13)], [
            'Internal Reference', 'Product', 'Location', 'Inventory Unit', 'Loose Count',
            'Packaging 1', 'Count 1', 'Packaging 2', 'Count 2',
            'Lot / Serial', 'Stock Package', 'Owner',
        ])
        self.assertEqual(book['_InventoryMetadata'].sheet_state, 'veryHidden')
        self.assertFalse(any(cell.data_type == 'f' for row in sheet for cell in row))
        self.assertTrue(sheet.column_dimensions['M'].hidden)
        self.assertIn('contains 250 Count ml', sheet['F2'].value)
        book.close()

    def test_audit_records_cannot_be_forged_by_rpc(self):
        template = self._template().with_context(_tha_packaging_operation=True)
        with self.assertRaises(AccessError):
            template.write({'snapshot': {}})
        with self.assertRaises(AccessError):
            self.templates.with_context(_tha_packaging_operation=True).create({
                'name': 'Fake', 'token': 'fake', 'company_id': self.env.company.id, 'snapshot': {},
            })
        batch = self._preview(template, {self.quants[0].id: {5: 1}}).with_context(_tha_packaging_operation=None)
        with self.assertRaises(AccessError):
            batch.line_ids.write({'new_count': 999})

    def test_cross_company_template(self):
        template = self._template()
        data = self._edit_file(template, {self.quants[0].id: {5: 1}})
        company = self.env['res.company'].create({'name': 'Packaging Count Other Company'})
        with self.assertRaises(UserError):
            self.batches.with_company(company)._preview_upload(data, 'Counts.xlsx')

    def test_upload_wizard_and_physical_inventory_actions(self):
        download = self.quants.tha_action_download_packaging_counts()
        self.assertEqual(download['type'], 'ir.actions.act_url')
        self.assertIn('file_data', download['url'])
        action = self.env['stock.quant'].tha_action_upload_packaging_counts()
        self.assertEqual(action['res_model'], 'tha.inventory.count.upload')
        # List buttons pass selected IDs, including an empty list, through RPC.
        for selected_ids in ([], self.quants.ids):
            rpc_action = call_kw(self.env['stock.quant'], 'tha_action_upload_packaging_counts',
                                 [selected_ids], {})
            self.assertEqual(rpc_action['res_model'], 'tha.inventory.count.upload')
            self.assertEqual(rpc_action['target'], 'new')
        template = self._template()
        wizard = self.env['tha.inventory.count.upload'].create({
            'file_data': base64.b64encode(self._edit_file(template, {self.quants[0].id: {5: 5}})),
            'filename': 'Counts.xlsx',
        })
        preview = wizard.action_preview()
        batch = self.batches.browse(preview['res_id'])
        self.assertEqual(batch.line_ids.new_count, 5)
        self.assertEqual(batch.action_view_inventory()['domain'], [('id', 'in', [self.quants[0].id])])

    def test_stock_user_cannot_update_counts_with_packaging_actions(self):
        user = self.env['res.users'].create({
            'name': 'Packaging Count Stock User', 'login': 'packaging_count_stock_user',
            'group_ids': [Command.set([self.env.ref('stock.group_stock_user').id])],
            'company_id': self.env.company.id, 'company_ids': [Command.set([self.env.company.id])],
        })
        with self.assertRaises(AccessError):
            self.quants.with_user(user).tha_action_download_packaging_counts()
        with self.assertRaises(AccessError):
            self.env['stock.quant'].with_user(user).tha_action_upload_packaging_counts()

    def test_dimension_changed_after_export(self):
        template = self._template()
        self.quants[0].location_id = self.location.copy({'name': 'Changed Count Location'})
        with self.assertRaisesRegex(UserError, 'changed product, location'):
            self._preview(template, {self.quants[0].id: {5: 1}})

    def test_950_products_scale_and_new_product_export(self):
        products = self.env['product.product'].create([
            {'name': 'Count Product %04d' % index, 'default_code': 'COUNT-%04d' % index,
             'is_storable': True, 'uom_id': self.unit.id,
             'uom_ids': [Command.set([self.bottle.id, self.box.id] if index % 2 else [])]}
            for index in range(950)
        ])
        quants = self.env['stock.quant'].create([
            {'product_id': product.id, 'location_id': self.location.id, 'quantity': 5} for product in products
        ])
        template = self._template(quants)
        self.assertEqual(len(template.snapshot['rows']), 950)
        self.assertEqual(template.snapshot['max_slots'], 2)
        data = self._edit_file(template, {quant.id: {5: 2} for quant in quants})
        count_before = self.env['stock.quant'].search_count([])
        batch = self.batches._preview_upload(data, '950 Products.xlsx')
        self.assertEqual(len(batch.line_ids), 950)
        batch.action_confirm()
        self.assertEqual(set(quants.mapped('inventory_quantity')), {2})
        self.assertEqual(set(quants.mapped('quantity')), {5})
        self.assertEqual(self.env['stock.quant'].search_count([]), count_before)
        new_product = self.env['product.product'].create({'name': 'New Count Product', 'is_storable': True, 'uom_id': self.unit.id})
        new_quant = self.env['stock.quant'].create({'product_id': new_product.id, 'location_id': self.location.id})
        fresh = self._template(quants | new_quant)
        self.assertIn(str(new_quant.id), fresh.snapshot['rows'])
